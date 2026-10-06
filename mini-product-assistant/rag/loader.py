"""Carga y validación de productos desde Excel."""

from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from pydantic import ValidationError

from models import ProductRecord

REQUIRED_COLUMNS = {"product_id", "category", "product_name"}


# Representa errores de formato, lectura o validación del archivo de productos.
class ProductLoadError(ValueError):
    """El archivo no cumple el contrato de datos de la demo."""


# Convierte celdas vacías o strings en blanco a None y conserva el resto.
def _normalize_cell(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, str):
        cleaned = value.strip()
        return cleaned or None
    return value


# Lee la hoja Products y devuelve todas sus filas como modelos validados.
def load_products(path: str | Path) -> list[ProductRecord]:
    """Lee la hoja Products, normaliza vacíos y valida cada fila."""

    source = Path(path)
    if not source.is_file():
        raise ProductLoadError(f"No se encontró el archivo de productos: {source}")

    try:
        workbook = load_workbook(source, read_only=True, data_only=True)
    except Exception as error:
        raise ProductLoadError(f"No se pudo abrir el Excel: {source}") from error

    try:
        if "Products" not in workbook.sheetnames:
            raise ProductLoadError("El Excel no contiene la hoja requerida 'Products'.")

        sheet = workbook["Products"]
        rows = sheet.iter_rows(values_only=True)
        try:
            raw_headers = next(rows)
        except StopIteration as error:
            raise ProductLoadError("La hoja 'Products' está vacía.") from error

        headers = [str(value).strip() if value is not None else "" for value in raw_headers]
        missing = REQUIRED_COLUMNS - set(headers)
        if missing:
            names = ", ".join(sorted(missing))
            raise ProductLoadError(f"Faltan columnas requeridas: {names}.")

        products: list[ProductRecord] = []
        seen_ids: set[str] = set()
        for excel_row, values in enumerate(rows, start=2):
            normalized = [_normalize_cell(value) for value in values]
            if all(value is None for value in normalized):
                continue
            row = dict(zip(headers, normalized, strict=False))
            try:
                product = ProductRecord.model_validate(row)
            except ValidationError as error:
                raise ProductLoadError(
                    f"La fila {excel_row} de 'Products' no es válida: {error}"
                ) from error
            if product.product_id in seen_ids:
                raise ProductLoadError(
                    f"El product_id '{product.product_id}' está duplicado."
                )
            seen_ids.add(product.product_id)
            products.append(product)

        return products
    finally:
        workbook.close()
